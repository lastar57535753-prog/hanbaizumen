# -*- coding: utf-8 -*-
"""下帯に「囲い込み防止」帯（QR＋文言＋C21ロゴ／92×27mm）を入れる。

全物件に必ず入れる運用なので build_all.py から既定で呼ばれる（外すときは --no-kakoikomi）。
帯・連絡先(id314)・案内文(id318)で下帯の右半分を分け合う。そのままでは幅が足りないので、
連絡先と案内文を縮めて詰め、空いたところへ帯を置く。QR空枠(id229)は帯に置き換わるので消す。
案内文の「QRコードからお問合せ下さい」の行は帯と重複するので落とし、2行目も短い言い回しにする。

usage: python add_kakoikomi_band.py <図面.pptx> [--band 別の帯.png]
"""
import os
import sys

from pptx import Presentation
from pptx.util import Cm, Pt

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_BAND = os.path.join(HERE, "..", "assets", "kakoikomi_band.png")

RATIO = 92.0 / 27.0      # 帯の原寸（mm）
BAND_W = 6.35            # 下帯の残り幅に収まる帯の幅(cm)
NOTE_ID, QRBOX_ID = 318, 229
FOOTER_ID, CONTACT_ID, TORIHIKI_ID = 146, 314, 319

CONTACT_BOX = (11.20, 4.60)              # 連絡先 x, 幅(cm)
CONTACT_PT = 8.5
NOTE_BOX = (16.00, 3.05)                 # 案内文 x, 幅(cm)
NOTE_PT = 7.0
NOTE_REWRITE = {"内覧依頼/資料請求/広告掲載依頼": "内覧・資料請求は担当まで"}

args = sys.argv[1:]
band = DEFAULT_BAND
if "--band" in args:
    i = args.index("--band")
    band = args[i + 1]
    del args[i:i + 2]
if not args:
    print(__doc__); sys.exit(1)
path = args[0]
if not os.path.exists(band):
    sys.exit(f"!! 囲い込み防止帯の画像がありません: {band}")

prs = Presentation(path)
slide = prs.slides[0]
byid = {sh.shape_id: sh for sh in slide.shapes}

if QRBOX_ID in byid:                            # 帯に置き換わるQR空枠は外す
    byid[QRBOX_ID]._element.getparent().remove(byid[QRBOX_ID]._element)

if CONTACT_ID in byid:                          # 連絡先を縮めて帯の場所をあける
    c = byid[CONTACT_ID]
    c.left, c.width = Cm(CONTACT_BOX[0]), Cm(CONTACT_BOX[1])
    for para in c.text_frame.paragraphs:
        for r in para.runs:
            if r.font.size is None or r.font.size > Pt(CONTACT_PT):
                r.font.size = Pt(CONTACT_PT)

if NOTE_ID in byid:                             # 案内文は縮めて連絡先の右へ
    note = byid[NOTE_ID]
    tf = note.text_frame
    for para in list(tf.paragraphs):
        if "QR" in para.text:                   # QRの案内は帯と重複するので落とす
            para._p.getparent().remove(para._p)
    for para in tf.paragraphs:
        new_text = NOTE_REWRITE.get(para.text)      # 段落は複数runに割れているので段落単位で見る
        if new_text:
            runs = para._p.findall(f"{A}r")
            for r in runs[1:]:
                para._p.remove(r)
            runs[0].find(f"{A}t").text = new_text
        for r in para.runs:
            r.font.size = Pt(NOTE_PT)
    note.left, note.width = Cm(NOTE_BOX[0]), Cm(NOTE_BOX[1])

w, h = Cm(BAND_W), Cm(BAND_W / RATIO)
right_edge = byid[TORIHIKI_ID].left if TORIHIKI_ID in byid else Cm(25.60)
x = right_edge - w - Cm(0.05)                              # 取引態様の左にぴたりと寄せる
if FOOTER_ID in byid:                                       # 下帯の中で上下中央に
    f = byid[FOOTER_ID]
    y = f.top + (f.height - h) // 2
else:
    y = Cm(18.95)

slide.shapes.add_picture(band, x, y, w, h)
prs.save(path)
print(f"連絡先(id{CONTACT_ID})を幅{CONTACT_BOX[1]}cm/{CONTACT_PT}pt、"
      f"案内文(id{NOTE_ID})を x{NOTE_BOX[0]}cm/{NOTE_PT}pt に詰めた")
print(f"囲い込み防止帯を配置: {os.path.basename(band)} "
      f"（{w / 360000:.2f}×{h / 360000:.2f}cm ／ x={x / 360000:.2f} y={y / 360000:.2f}）")
