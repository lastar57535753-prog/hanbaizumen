# -*- coding: utf-8 -*-
"""マンション用テンプレから土地用テンプレを作る。
意匠（座標・書体・色・帯・バッジ枠）は一切触らず、物件概要表のラベルと
プレースホルダ名だけを土地用に差し替える。行数23はそのままなので高さも変わらない。"""
import sys
from pptx import Presentation

ROWS = [
    ("所 在 地", "ADDRESS"),
    ("権　利", "RIGHT"),
    # ラベルは「敷地面積」。「土地面積」だと LibreOffice が先頭2文字「土地」を入れ替えて
    # 重ねて描く不具合があり、PDF で「地面積」に見える（PowerPoint では正常）。
    ("敷地面積", "LAND_AREA"),
    ("坪 単 価", "TSUBO_PRICE"),
    ("地　目", "CHIMOKU"),
    ("地　勢", "TOPOGRAPHY"),
    ("用途地域", "ZONING"),
    ("建ぺい率", "BUILDING_RATIO"),
    ("容 積 率", "FLOOR_RATIO"),
    ("防火指定", "FIRE_ZONE"),
    ("高度地区", "HEIGHT_ZONE"),
    ("都市計画", "CITY_PLAN"),
    ("接道方位", "ROAD_DIR"),
    ("接道幅員", "ROAD_WIDTH"),
    ("接面距離", "ROAD_FRONT"),
    ("私道負担", "PRIVATE_ROAD"),
    ("セットバック", "SETBACK"),
    ("建築条件", "BUILD_COND"),
    ("上下水道", "WATER_SEWER"),
    ("ガ　ス", "GAS"),
    ("特　記", "SPECIAL"),
    ("現　況", "STATUS"),
    ("引　渡", "HANDOVER"),
]


def set_cell(cell, text):
    """1つ目のrunの書式を保ったまま文字だけ差し替える。"""
    p = cell.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for extra in p.runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        p.add_run().text = text


def set_shape_text(sh, lines):
    """段落ごとに1つ目のrunだけ残して差し替える（段落数が一致する前提）。"""
    tf = sh.text_frame
    for para, text in zip(tf.paragraphs, lines):
        if para.runs:
            para.runs[0].text = text
            for extra in para.runs[1:]:
                extra._r.getparent().remove(extra._r)


def main(src, dst):
    prs = Presentation(src)
    slide = prs.slides[0]
    byid = {}

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
            else:
                byid[sh.shape_id] = sh
    walk(slide.shapes)

    tbl = byid[307].table
    assert len(tbl.rows) == len(ROWS), (len(tbl.rows), len(ROWS))
    for r, (label, key) in enumerate(ROWS):
        set_cell(tbl.cell(r, 0), label)
        set_cell(tbl.cell(r, 1), "{{%s}}" % key)
    print(f"物件概要表({307}) を土地用 {len(ROWS)} 行に差し替え")

    # 左上の縦書きラベル「マンション」→「土　地」
    set_shape_text(byid[276], ["土　地"])
    print("種別ラベル(276) を「土　地」に差し替え")

    # 間取図脇の小表示（間取り/向き・面積・階）→ 坪数・㎡・建ぺい/容積
    set_shape_text(byid[28], ["{{LAND_TSUBO}}", "{{AREA_MINI}}", "{{BUILDING_RATIO}}/{{FLOOR_RATIO}}"])
    print("小表示(28) を 坪数・面積・建ぺい率/容積率 に差し替え")

    prs.save(dst)
    print("saved:", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
