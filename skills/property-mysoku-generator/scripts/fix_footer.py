# -*- coding: utf-8 -*-
"""下帯（会社情報）の文字が枠からはみ出して折り返すのを直す。テンプレを直接書き換える。

  usage: python fix_footer.py assets/template_mansion.pptx assets/template_land.pptx

直す内容（実測して決めた値）:
  id=176 2行目（免許・協会）  7.0pt → 6.5pt   9.38cm あって枠(9.26cm)に入らなかった
  id=314 TEL/携帯/Mail/定休  11pt  → 8.5pt   Mail が 6.01cm、定休が 7.18cm で折り返していた
         幅 6.20 → 6.30cm（右隣 318 の手前まで）
  id=318 広告掲載の3行        10.5pt → 10pt   x 17.06 → 17.55、幅 6.88 → 6.00cm
"""
import sys, os
from pptx import Presentation
from pptx.util import Pt, Cm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def fix(path):
    prs = Presentation(path)
    byid = {sh.shape_id: sh for sh in walk(prs.slides[0].shapes)}

    sh = byid.get(176)
    if sh is not None and sh.has_text_frame:
        paras = sh.text_frame.paragraphs
        if len(paras) >= 2:
            for r in paras[1].runs:
                r.font.size = Pt(6.5)

    sh = byid.get(314)
    if sh is not None and sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8.5)
        sh.width = Cm(6.30)

    sh = byid.get(318)
    if sh is not None and sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
        sh.left = Cm(17.55)
        sh.width = Cm(6.00)

    prs.save(path)
    print("下帯を調整:", os.path.basename(path))


if __name__ == "__main__":
    for p in sys.argv[1:]:
        fix(p)
